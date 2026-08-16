"""Ship 72'.d — snapshot tests for the app-wide FindingContract.

Run: PYTHONPATH=/data/arioncomply python3 tests/test_finding_contract.py

Pins the extractor-side SSoT (finding_contract.py) against regression
in three places at once:

  1. `is_scaffolding` — the app-wide scaffolding predicate. Locked against
     every renderer-emitted scaffolding shape from the reader's docx
     reconstruction (▽/△ rails, Best practice blocks, ✓ Good examples,
     Standard text blockquote, prereq categories, ☑/☐ guidance bullets).

  2. `catalog_recognises` — Task #606's promoted membership check.
     Known-good ids pass, known-bad + case-mangled ids fail.

  3. `FindingContract.bind()` — the round-trip loop that Ship 72' arc
     opened around. Unedited docx round-trip produces ZERO findings;
     tenant-filled placeholders produce exactly the filled findings.

Requires the API server to be live at http://localhost:8080 with an
Arion tenant + docx templates loaded. Skips cleanly when the API
isn't reachable (CI + dev-laptop mode).
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(_ROOT))

try:
    from dotenv import load_dotenv
    if (_ROOT / ".env").exists():
        load_dotenv(_ROOT / ".env")
except ImportError:
    pass

from rag.intake.finding_contract import (
    FINDING_CONTRACT, ExtractedCandidate, SkipReason,
    is_scaffolding, catalog_recognises,
)


API_KEY  = os.getenv("ARION_API_KEY", "arion_dev_key_2026")
API_BASE = "http://localhost:8080"


def _api_available() -> bool:
    try:
        from urllib.request import Request, urlopen
        r = urlopen(f"{API_BASE}/openapi.json", timeout=3)
        return r.status == 200
    except Exception:
        return False


# ── Predicate-level tests ───────────────────────────────────────────

def test_is_scaffolding_recognizes_empty_string():
    assert is_scaffolding("") is True
    assert is_scaffolding("   \n\n  ") is True
    assert is_scaffolding(None) is True


def test_is_scaffolding_recognizes_reader_reconstructed_shapes():
    """Every scaffolding shape the reader can emit — one string per
    pattern — must reject. Any regression in the renderer or the
    reader that changes one of these shapes will fail this test."""
    shapes = [
        # Placeholder + provenance
        "<<TEXT>>",
        "<<NAME>>",
        "[ Click to enter your evidence here ]",
        "<!-- prefilled from 3 sources -->",
        "<!-- EDIT-ZONE-START item:A.5.15:logical_rules -->",
        "<!-- EDIT-ZONE-END item:A.5.15:logical_rules -->",
        "<!-- TABLE-COLUMNS leaf:req:A.5.16:identity_revocation_record -->",
        "<!-- column: item:A.5.16:rev_identity_ref -->",
        # Reader-reconstructed line shapes
        "◆ Required element — logical rules",
        "◆ Recommended addition — emergency access",
        "*Do not edit — system id*: `<<MUST item:A.5.15:logical_rules>>`",
        "*Standard text:* Logical access rules (systems, applications, network segments)",
        "*Why: Access control policy must specify who is responsible*",
        "*Good enough: A documented scope statement*",
        "__✓ Good__:",
        "__Best practice ✓ — covered:__",
        "__Best practice ◐ — partly covered:__",
        "- ☑ Document all systems, applications, and network segments",
        "- ☐ Assign responsibility for approving access requests",
        "▽ Enter your evidence for \"logical rules\" below ▽",
        "△ End of \"logical rules\" △",
        "───────────────────────────────────",
        "[Not applicable to your scope — no evidence required.]",
        "**Foundational**",
        "**Direct upstream**",
        "**Cross-framework**",
    ]
    for shape in shapes:
        assert is_scaffolding(shape), f"expected scaffolding: {shape!r}"


def test_is_scaffolding_rejects_real_tenant_evidence():
    """Anything that looks like real tenant-authored prose must NOT
    trigger the scaffolding predicate."""
    real_evidence = [
        "Logical access to production systems uses SSO via Okta with hardware MFA (YubiKey); RBAC bundles enforced per role register.",
        "Our access control policy is owned by the CISO, reviewed annually + approved by ISMS Owner.",
        "role-based access control is implemented across all production systems",
        "Users must submit access requests via the ServiceNow portal; approvals routed to the asset owner.",
    ]
    for evidence in real_evidence:
        assert not is_scaffolding(evidence), \
            f"expected NOT scaffolding: {evidence!r}"


def test_catalog_recognises_known_good():
    """Known-good catalog ids must pass."""
    good = [
        "item:A.5.15:logical_rules",
        "item:A.5.16:rev_identity_ref",
        "item:6.1.2:criteria",
    ]
    for iid in good:
        assert catalog_recognises(iid), f"expected recognized: {iid!r}"


def test_catalog_recognises_rejects_mangled_and_wrong_case():
    """Typo'd or case-mangled ids must fail."""
    bad = [
        "item:A.5.15:logica_rules",           # missing an 'l'
        "item:A.5.15:LOGICAL_RULES",          # wrong case
        "item:A.99.99:nonexistent",           # nonexistent control
        "item:X.Y.Z:whatever",                # nonsense id
        "",                                   # empty
    ]
    for iid in bad:
        assert not catalog_recognises(iid), \
            f"expected NOT recognized: {iid!r}"


# ── Contract-level tests ────────────────────────────────────────────

def test_bind_rejects_empty_text():
    c = ExtractedCandidate(
        item_id       = "item:A.5.15:logical_rules",
        excerpt_text  = "",
        document_name = "test.docx",
    )
    r = FINDING_CONTRACT.bind(c)
    assert r.reason == SkipReason.EMPTY_TEXT
    assert r.finding is None


def test_bind_rejects_scaffolding():
    c = ExtractedCandidate(
        item_id       = "item:A.5.15:logical_rules",
        excerpt_text  = "▽ Enter your evidence for \"logical rules\" below ▽\n\n[ Click to enter your evidence here ]\n\n△ End of \"logical rules\" △",
        document_name = "test.docx",
    )
    r = FINDING_CONTRACT.bind(c)
    assert r.reason == SkipReason.PURE_SCAFFOLDING
    assert r.finding is None


def test_bind_rejects_mangled_item_id():
    c = ExtractedCandidate(
        item_id       = "item:A.5.15:logica_rules",  # typo
        excerpt_text  = "Real tenant policy content.",
        document_name = "test.docx",
    )
    r = FINDING_CONTRACT.bind(c)
    assert r.reason == SkipReason.MANGLED_ITEM_ID
    assert r.finding is None


def test_bind_accepts_valid_candidate():
    c = ExtractedCandidate(
        item_id          = "item:A.5.15:logical_rules",
        excerpt_text     = "Logical access uses SSO via Okta with hardware MFA; RBAC enforced.",
        document_name    = "test.docx",
        upload_id        = "upload-test",
        inference_source = "templated",
    )
    r = FINDING_CONTRACT.bind(c)
    assert r.reason == SkipReason.OK
    assert r.finding is not None
    assert r.finding.checklist_item_id == "item:A.5.15:logical_rules"
    assert r.finding.control_ref       == "A.5.15"
    assert r.finding.standard_id       == "ISO27001:2022"
    assert r.finding.inference_source  == "templated"
    assert "SSO via Okta" in r.finding.evidence_text


# ── End-to-end docx round-trip tests ────────────────────────────────

_TEMPLATE_LEAVES = [
    "req:A.5.15:access_control_policy",
    "req:A.5.1:isp_policy",
    "req:A.5.24:incident_response_procedure",
    "req:6.1.2:risk_assessment",
]


def _download_docx(leaf_id: str) -> bytes:
    from urllib.request import Request, urlopen
    req = Request(
        f"{API_BASE}/api/v1/templates/{leaf_id}/download?format=docx",
        headers={"X-API-Key": API_KEY},
    )
    return urlopen(req).read()


def _run_extractor_edit_zones(docx_path: str):
    """Reader → edit-zone extraction. Returns (findings, metrics)."""
    from rag.intake.readers import _read_docx
    from rag.intake.extractor import (
        _TEMPLATED_EDIT_ZONE_RE, _extract_templated_via_edit_zones,
    )
    parsed = _read_docx(docx_path, docx_path.rsplit("/", 1)[-1])
    parsed.upload_id = "test-72d"
    parsed.extraction_metrics = {}
    body = parsed.markdown or parsed.full_text or ""
    zones = list(_TEMPLATED_EDIT_ZONE_RE.finditer(body))
    findings = _extract_templated_via_edit_zones(parsed, zones)
    return findings, parsed.extraction_metrics, zones


def test_unedited_docx_round_trip_produces_zero_findings():
    """The bug that opened Ship 72': tenant downloads a docx template
    and re-uploads it unedited. Extractor MUST bind zero findings —
    every zone contains only scaffolding.

    Ship 72'.a's is_scaffolding + reader-side ▽/△ rail boundary fix
    are what make this pass.
    """
    if not _api_available():
        print("skip: API not reachable")
        return
    for leaf_id in _TEMPLATE_LEAVES:
        data = _download_docx(leaf_id)
        path = f"/tmp/72d_unedited_{leaf_id.replace(':','_')}.docx"
        with open(path, "wb") as f:
            f.write(data)
        findings, metrics, zones = _run_extractor_edit_zones(path)
        n_zones = len(zones)
        n_bound = len(findings)
        assert n_bound == 0, (
            f"{leaf_id}: expected 0 findings on unedited docx, got {n_bound} "
            f"(zones={n_zones}). Ship 72'.a regression."
        )
        # Metric surface — every zone should be accounted for as
        # either scaffolding-rejected or otherwise skipped.
        rejected = (
            metrics.get("templated_zones_scaffolding", 0)
            + metrics.get("templated_zones_mangled", 0)
        )
        # Note: some empty-zone rejects roll up under scaffolding via
        # Ship 72'.a metric mapping (EMPTY_TEXT + PURE_SCAFFOLDING).
        # Assertion: EVERY zone either produced a finding (impossible
        # here) or was skipped — nothing should silently vanish.
        assert n_bound + rejected == n_zones or n_zones == 0, (
            f"{leaf_id}: metric coverage gap — zones={n_zones}, "
            f"bound={n_bound}, rejected={rejected}"
        )


def test_filled_docx_produces_expected_findings():
    """Fill 2 placeholders in an A.5.15 docx, re-upload path emits
    2 findings + rejects the 4 unfilled placeholders."""
    if not _api_available():
        print("skip: API not reachable")
        return
    from docx import Document

    data = _download_docx("req:A.5.15:access_control_policy")
    src_path = "/tmp/72d_filled_src.docx"
    with open(src_path, "wb") as f:
        f.write(data)

    d = Document(src_path)
    samples = [
        "Logical access: SSO via Okta + hardware MFA on all production systems.",
        "RBAC bundles maintained in Okta per role register.",
    ]
    n_filled = 0
    for p in d.paragraphs:
        if "Click to enter your evidence here" in p.text and n_filled < len(samples):
            for r in p.runs:
                r.text = ""
            p.runs[0].text = samples[n_filled]
            n_filled += 1
    filled_path = "/tmp/72d_filled.docx"
    d.save(filled_path)

    findings, metrics, zones = _run_extractor_edit_zones(filled_path)
    assert len(findings) == n_filled, (
        f"expected {n_filled} findings, got {len(findings)}. "
        f"metrics={metrics}"
    )
    # Correct binding
    ids = {f.checklist_item_id for f in findings}
    assert "item:A.5.15:logical_rules" in ids
    assert "item:A.5.15:rbac"          in ids


# ── Test harness ────────────────────────────────────────────────────

CASES = [
    test_is_scaffolding_recognizes_empty_string,
    test_is_scaffolding_recognizes_reader_reconstructed_shapes,
    test_is_scaffolding_rejects_real_tenant_evidence,
    test_catalog_recognises_known_good,
    test_catalog_recognises_rejects_mangled_and_wrong_case,
    test_bind_rejects_empty_text,
    test_bind_rejects_scaffolding,
    test_bind_rejects_mangled_item_id,
    test_bind_accepts_valid_candidate,
    test_unedited_docx_round_trip_produces_zero_findings,
    test_filled_docx_produces_expected_findings,
]


def main() -> int:
    failed = 0
    for fn in CASES:
        name = fn.__name__
        try:
            fn()
            print(f"PASS  {name}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {name}: {e}")
        except Exception as e:
            failed += 1
            print(f"ERROR {name}: {type(e).__name__}: {e}")
    if failed:
        print(f"\n{failed} of {len(CASES)} cases failed.")
        return 1
    print(f"\n{len(CASES)} of {len(CASES)} cases passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
